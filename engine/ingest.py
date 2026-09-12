import re
from dataclasses import dataclass, field
from pathlib import Path
from rapidfuzz import fuzz

SECTION_LEXICON = {
    "EXPERIENCE": ["experience", "work experience", "professional experience", "employment history", "work ex", "employment", "internships", "internship experience", "work history"],
    "PROJECTS": ["projects", "personal projects", "key projects", "portfolio", "academic projects", "notable projects"],
    "SKILLS": ["skills", "technical skills", "technologies", "tech stack", "core competencies", "tools", "languages and tools", "technical proficiency"],
    "EDUCATION": ["education", "academic background", "qualifications", "academics", "academic qualifications"],
    "OTHER": ["certifications", "achievements", "awards", "summary", "objective", "about", "profile", "interests", "hobbies", "extracurricular", "positions of responsibility", "activities", "languages"],
}
SECTION_WEIGHT = {"EXPERIENCE": 1.0, "PROJECTS": 0.85, "SKILLS": 0.6, "EDUCATION": 0.5, "OTHER": 0.5, "HEADER": 0.5}

MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
DATE_RANGE = re.compile(
    rf"((?:{MONTHS})[a-z]*\.?\s*'?\d{{2,4}}|\d{{1,2}}/\d{{2,4}}|\b(?:19|20)\d{{2}}\b)\s*(?:-|–|—|to|until)\s*"
    rf"((?:{MONTHS})[a-z]*\.?\s*'?\d{{2,4}}|\d{{1,2}}/\d{{2,4}}|\b(?:19|20)\d{{2}}\b|present|current|now|ongoing|\d{{2}}\b)",
    re.I,
)
SHORT_RANGE = re.compile(rf"\b((?:{MONTHS})[a-z]*)\.?\s*(?:-|\u2013|\u2014|to)\s*((?:{MONTHS})[a-z]*|present|current|now)\.?\s*'?(\d{{2,4}})\b", re.I)
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
PHONE = re.compile(r"\+?\(?\d[\d\s\-()]{7,14}\d")


def _is_phone(m):
    digits = re.sub(r"\D", "", m.group(0))
    return 10 <= len(digits) <= 13 and not re.search(r"(?:19|20)\d{2}\s*[-–]\s*(?:19|20)?\d{2}", m.group(0))
URL = re.compile(r"(https?://\S+|www\.\S+|linkedin\.com/\S+|github\.com/\S+)", re.I)
BULLET = re.compile(r"^\s*[•\-\*▪◦●·»➢✓–—]\s*")


@dataclass
class Chunk:
    chunk_id: str
    resume_id: str
    section: str
    text: str
    line_no: int
    weight: float
    depth: str = "listed"
    demonstrated: bool = False


@dataclass
class Resume:
    resume_id: str
    name: str
    filename: str
    raw_text: str
    lines: list
    chunks: list = field(default_factory=list)
    sections: dict = field(default_factory=dict)
    parse_method: str = "text"
    parse_confidence: float = 1.0
    fuzzy_headers: int = 0
    dates_found: int = 0
    months_experience: int = 0
    pii: dict = field(default_factory=dict)
    hidden_chars: int = 0
    injected: list = field(default_factory=list)
    integrity: list = field(default_factory=list)
    duplicates: list = field(default_factory=list)


CID = re.compile(r"\(cid:\d+\)")


def clean_text(text: str):
    text = CID.sub("• ", text)
    text = text.replace("", "• ").replace("‣", "• ").replace("▪", "• ")
    return text


def extract_text(path: Path):
    text, method = _extract_text(path)
    return clean_text(text), method


def _extract_text(path: Path):
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore"), "text"
    if suffix == ".docx":
        return _docx_text(path), "docx"
    if suffix == ".xml":
        return _xml_text(path), "xml"
    text = ""
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            text = "\n".join((p.extract_text() or "") for p in pdf.pages)
            HIDDEN[str(path)] = _hidden_chars(pdf)
        if len(text.strip()) > 200:
            return text, "pdfplumber"
    except Exception:
        pass
    try:
        import fitz
        doc = fitz.open(str(path))
        text2 = "\n".join(page.get_text() for page in doc)
        if len(text2.strip()) > len(text.strip()):
            return text2, "pymupdf"
    except Exception:
        pass
    if text.strip():
        return text, "pdfplumber"
    return "", "failed"


HIDDEN = {}


def _hidden_chars(pdf):
    """Characters a human would not see: white/near-white fill or sub-4pt size."""
    n = 0
    for page in pdf.pages:
        for ch in page.chars:
            col = ch.get("non_stroking_color")
            white = False
            if isinstance(col, (list, tuple)) and col:
                try:
                    white = all(float(v) >= 0.93 for v in col)
                except Exception:
                    white = False
            elif isinstance(col, (int, float)):
                white = float(col) >= 0.93
            if white or (ch.get("size") or 10) < 4:
                if not (ch.get("text") or "").isspace():
                    n += 1
    return n


def _docx_text(path: Path):
    try:
        import docx
        d = docx.Document(str(path))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        return "\n".join(lines)
    except Exception:
        import zipfile
        with zipfile.ZipFile(str(path)) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
        xml = re.sub(r"</w:p>", "\n", xml)
        return re.sub(r"<[^>]+>", "", xml)


def _xml_text(path: Path):
    """Structured XML resumes: emit one line per leaf, promoting section-like tags to headers."""
    import xml.etree.ElementTree as ET
    HEADERS = {"education": "EDUCATION", "experience": "EXPERIENCE", "workexperience": "EXPERIENCE", "internships": "EXPERIENCE", "projects": "PROJECTS", "technicalskills": "SKILLS", "skills": "SKILLS", "certifications": "CERTIFICATIONS", "achievements": "ACHIEVEMENTS"}
    try:
        root = ET.fromstring(path.read_bytes())
    except Exception:
        return re.sub(r"<[^>]+>", "\n", path.read_text(encoding="utf-8", errors="ignore"))
    out = []

    def walk(el, depth):
        tag = el.tag.lower()
        if tag in HEADERS and depth <= 2:
            out.append("")
            out.append(HEADERS[tag])
        text = (el.text or "").strip()
        kids = list(el)
        if text and not kids:
            prefix = "- " if depth >= 3 else ""
            out.append(prefix + " ".join(text.split()))
        for k in kids:
            walk(k, depth + 1)
            tail = (k.tail or "").strip()
            if tail:
                out.append(tail)

    walk(root, 0)
    return "\n".join(out)


def classify_header(line: str):
    s = line.strip().rstrip(":").strip().lower()
    if not s or len(s) > 40 or len(s.split()) > 5:
        return None, False
    best, best_score, exact = None, 0, False
    for section, names in SECTION_LEXICON.items():
        for n in names:
            if s == n:
                return section, False
            sc = fuzz.ratio(s, n)
            if sc > best_score:
                best, best_score = section, sc
    if best_score >= 80:
        return best, True
    return None, False


def normalize_dates(text: str):
    months = 0
    found = 0
    for m in SHORT_RANGE.finditer(text):
        found += 1
        try:
            y = _year(m.group(3))
            a, b = _ym(f"{m.group(1)} {y}"), _ym(f"{m.group(2)} {y}")
            if a and b and b >= a:
                months += max(1, b - a + 1)
        except Exception:
            pass
    text = SHORT_RANGE.sub(" ", text)
    for m in DATE_RANGE.finditer(text):
        found += 1
        try:
            a, b = _ym(m.group(1)), _ym(m.group(2))
            b = min(b, NOW) if b else b
            if a and b and b >= a:
                months += max(1, b - a)
        except Exception:
            pass
    return found, min(months, 120)


NOW = 2026 * 12 + 8
MONTH_IDX = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}


def _ym(s):
    """Absolute month index (year*12+month); unknown month defaults to mid-year."""
    y = _year(s)
    if not y:
        return None
    low = s.lower()
    mo = next((MONTH_IDX[k] for k in MONTH_IDX if k in low), None)
    if mo is None:
        mm = re.match(r"(\d{1,2})/", low)
        mo = int(mm.group(1)) - 1 if mm else (8 if any(w in low for w in ("present", "current", "now", "ongoing")) else 5)
    return y * 12 + mo


def _year(s, ref=None):
    s = s.lower()
    if s in ("present", "current", "now", "ongoing"):
        return 2026
    m = re.search(r"(?:19|20)\d{2}", s)
    if m:
        return int(m.group(0))
    m = re.search(r"'?(\d{2})\b", s)
    if m:
        v = int(m.group(1))
        if ref and v < 100:
            return 2000 + v if v < 50 else 1900 + v
        return 2000 + v if v < 50 else 1900 + v
    return None


def guess_name(lines):
    for line in lines[:6]:
        s = line.strip()
        if not s or EMAIL.search(s) or any(_is_phone(m) for m in PHONE.finditer(s)) or URL.search(s):
            continue
        if 1 < len(s.split()) <= 4 and len(s) < 40 and not any(ch.isdigit() for ch in s):
            return s.title() if s.isupper() else s
    return "Unknown Candidate"


def strip_pii(text: str, name: str):
    t = EMAIL.sub("[email]", text)
    t = PHONE.sub(lambda m: "[phone]" if _is_phone(m) else m.group(0), t)
    t = URL.sub("[url]", t)
    if name and name != "Unknown Candidate":
        t = re.sub(re.escape(name), "[name]", t, flags=re.I)
    t = re.sub(r"\b(he|she|his|her|him|hers)\b", "they", t, flags=re.I)
    return t


def parse_resume(path: Path, resume_id: str) -> Resume:
    raw, method = extract_text(path)
    lines = [l.rstrip() for l in raw.splitlines()]
    name = guess_name(lines)
    r = Resume(resume_id=resume_id, name=name, filename=path.name, raw_text=raw, lines=lines, parse_method=method)

    section = "HEADER"
    buffer, buffer_start = [], 0
    chunk_idx = 0

    def flush():
        nonlocal buffer, chunk_idx
        if buffer:
            text = " ".join(x.strip() for x in buffer).strip()
            text = re.sub(r"\s+", " ", text)
            if len(text) > 12:
                r.chunks.append(Chunk(f"{resume_id}_c{chunk_idx:03d}", resume_id, section, text, buffer_start, SECTION_WEIGHT.get(section, 0.5)))
                chunk_idx += 1
            buffer = []

    for i, line in enumerate(lines):
        sec, fuzzy = classify_header(line)
        if sec:
            flush()
            section = sec
            r.sections.setdefault(sec, i)
            if fuzzy:
                r.fuzzy_headers += 1
            continue
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if BULLET.match(line) or (section == "SKILLS" and len(stripped) < 120):
            flush()
            buffer = [BULLET.sub("", line)]
            buffer_start = i
        elif buffer and len(" ".join(buffer)) < 260 and not stripped[0].isupper():
            buffer.append(stripped)
        else:
            flush()
            buffer = [stripped]
            buffer_start = i
    flush()

    exp_text = "\n".join(c.text for c in r.chunks if c.section == "EXPERIENCE")
    r.dates_found, r.months_experience = normalize_dates(exp_text if exp_text.strip() else "")
    from engine.integrity import assign_depth, strip_injections
    r.hidden_chars = HIDDEN.pop(str(path), 0)
    r.injected = strip_injections(r)
    assign_depth(r)
    known = sum(1 for s in ("EXPERIENCE", "PROJECTS", "SKILLS", "EDUCATION") if s in r.sections)
    r.parse_confidence = round(min(1.0, 0.45 + 0.12 * known + (0.07 if method != "failed" else -0.4) + min(0.1, len(r.chunks) / 200)), 2)
    r.pii = {"name": name, "emails": EMAIL.findall(raw)[:1], "blind_text": strip_pii(raw, name)}
    return r


def load_resumes(folder: Path):
    files = sorted([p for p in folder.iterdir() if p.suffix.lower() in (".pdf", ".docx", ".xml", ".txt", ".md")])
    rich = [p for p in files if p.suffix.lower() in (".pdf", ".docx", ".xml")]
    files = rich if rich else files
    return [parse_resume(p, f"r{i+1:02d}") for i, p in enumerate(files)]
